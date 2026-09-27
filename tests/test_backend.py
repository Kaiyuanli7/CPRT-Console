"""Backend tests - synthetic data only, no internet."""
from __future__ import annotations

import io
import sys
import time
import unittest
from unittest import mock

import pandas as pd

from helpers import FakeResp, TempData, copart_like_payload, iaa_like_payload, lots, split
from cprt import (capture, carriers, db, infer, jobs, macro, net, news, paths, report, scheduler, sec,
                  settings, tracker, weekly)


class SettingsTests(TempData, unittest.TestCase):
    def test_roundtrip_and_validation(self):
        s = settings.load()
        self.assertFalse(settings.is_configured(s))
        settings.update(user_name="Jane Doe", user_email="bad-email")
        self.assertFalse(settings.is_configured(settings.load()))
        s = settings.update(user_email="jane@school.edu", not_a_setting=1)
        self.assertEqual(settings.sec_user_agent(s), "Jane Doe jane@school.edu")
        self.assertNotIn("not_a_setting", settings.load())
        self.assertEqual(settings.schedule_text(s), "Mondays at 08:00")

    def test_odd_saved_series_never_break_loading(self):
        # load() runs before every page; a hand-edited settings.json must not take the console down
        for odd in ({"insurance_cpi": ["CUSR0000SETE"]}, {"insurance_cpi": "CUSR0000SETE"}, {"insurance_cpi": None},
                    {"insurance_cpi": ["a", "b", "c"]}, ["not", "a", "dict"]):
            settings.save({**settings.load(), "fred_series": odd})
            self.assertEqual(settings.load()["fred_series"], odd)

    def test_dead_fred_series_moves_to_bls(self):
        # FRED dropped the insurance index; saved settings that still name it are moved to the BLS copy.
        old = {**settings.DEFAULTS["fred_series"], "insurance_cpi": ["CUSR0000SETE", "Motor vehicle insurance prices"]}
        settings.save({**settings.load(), "fred_series": old})
        self.assertEqual(settings.load()["fred_series"]["insurance_cpi"][0], "BLS:CUUR0000SETE")


class InferenceTests(unittest.TestCase):
    def test_values_not_names(self):
        import random
        random.seed(3)
        recs = [{"a1": f"7{i:07d}", "zz": random.choice(["GEICO INSURANCE", "STATE FARM MUTUAL", "PROGRESSIVE CAS"]),
                 "q": random.choice(["TX - HOUSTON", "FL - MIAMI"]), "m": random.choice(["TOYOTA", "FORD"]),
                 "y": random.choice([2015, 2019, 2021]), "d": random.choice(["FRONT END", "HAIL"]),
                 "t": 1790000000000 + i, "odo": random.randint(900, 250000)} for i in range(40)]
        m, _ = infer.infer_mapping(recs)
        self.assertEqual(m, {"lot_id": "a1", "year": "y", "make": "m", "seller": "zz", "yard": "q",
                             "damage": "d", "sale_date": "t"})

    def test_override_wins(self):
        recs = [{"x": f"7{i:07d}", "s": "GEICO", "o": "USAA", "mk": "FORD", "yr": 2019} for i in range(10)]
        m, _ = infer.infer_mapping(recs, {"seller": "o"})
        self.assertEqual(m["seller"], "o")

    def test_state_parsing(self):
        self.assertEqual(infer.parse_state("", "Houston (TX)"), "TX")
        self.assertEqual(infer.parse_state("CA - SACRAMENTO"), "CA")
        self.assertEqual(infer.parse_state("Tampa, FL 33601"), "FL")
        self.assertEqual(infer.parse_state("nowhere"), "")

    def test_infinity_and_nan_do_not_crash(self):
        # Copart's make/model reference list has the code "INFINITY", and float("INFINITY") is infinite.
        recs = [{"code": v, "id": str(9000 + i)} for i, v in
                enumerate(["INFINITY", "Infinity", "-inf", "NaN", "1e999", "ACURA"] * 4)]
        recs.append({"code": float("inf"), "id": "9999"})
        self.assertEqual(infer.profile(recs)["code"]["year"], 0.0)
        self.assertFalse(infer._is_year(float("nan")))
        self.assertTrue(infer._is_year(2019) and infer._is_year("2019.0"))


class CaptureStorageTests(TempData, unittest.TestCase):
    def test_extract_store_and_fill_blanks(self):
        payloads = [{"url": "api", "page_url": "https://www.copart.com/s?st=TX", "data": copart_like_payload(12)},
                    {"url": "cfg", "page_url": "x", "data": {"flags": {"a": True}}}]
        rows = capture.extract(payloads, "copart")
        self.assertEqual(len(rows), 12)
        r = rows[0]
        self.assertEqual((r["state"], r["seller_raw"], r["make"], r["year"]), ("TX", "GEICO INSURANCE", "TOYOTA", "2019"))
        self.assertRegex(r["sale_date"], r"^\d{4}-\d{2}-\d{2}$")
        self.assertEqual(capture.find_total(payloads), 1500)
        # first sighting without seller, second the same day with seller -> blank gets filled
        capture.store([{**r, "seller_raw": ""}], "auto", "2026-09-28")
        capture.store([r], "auto", "2026-09-28")
        row = db.query("SELECT seller_raw FROM observations")[0]
        self.assertEqual(row["seller_raw"], "GEICO INSURANCE")
        self.assertEqual(db.scalar("SELECT COUNT(*) FROM observations"), 1)

    def test_iaa_shape_and_learning(self):
        payloads = [{"url": "api", "page_url": "https://www.iaai.com/Search?st=TX&utm_source=x#top",
                     "data": iaa_like_payload(15)}]
        rows = capture.extract(payloads, "iaa")
        self.assertEqual(len(rows), 15)
        self.assertEqual(rows[0]["state"], "TX")
        learned = capture.learn_sources(rows, "iaa", min_lots=10)
        self.assertEqual(len(learned), 1)
        self.assertEqual(learned[0]["url"], "https://www.iaai.com/Search?st=TX")  # tracking + fragment dropped
        self.assertEqual(learned[0]["label"], "TX")
        self.assertEqual(capture.learn_sources(rows[:5], "iaa", min_lots=10), [])

    def test_blocked_and_platform(self):
        self.assertTrue(capture.looks_blocked("Access Denied", ""))
        self.assertFalse(capture.looks_blocked("Search results", "lots"))
        self.assertEqual(capture.platform_of("https://www.iaai.com/x"), "iaa")
        self.assertEqual(capture.platform_of("https://www.copart.com/x"), "copart")

    def test_run_auto_with_fake_browser(self):
        db.upsert("sources", [{"platform": "copart", "url": "https://www.copart.com/a", "label": "TX", "enabled": 1},
                              {"platform": "iaa", "url": "https://www.iaai.com/blocked", "label": "TX", "enabled": 1},
                              {"platform": "iaa", "url": "https://www.iaai.com/b", "label": "TX", "enabled": 1}], ["url"])

        def fake(url, show, wait, scrolls):
            if "blocked" in url:
                return {"payloads": [], "title": "Pardon Our Interruption", "text": ""}
            data = copart_like_payload(20) if "copart" in url else iaa_like_payload(20)
            return {"payloads": [{"url": "api", "page_url": url, "data": data}], "title": "ok", "text": ""}

        s = settings.update(pause_between_pages=[0, 0])
        with mock.patch.object(net, "robots_allows", return_value=True):
            out = capture.run_auto(s, progress=lambda *a, **k: None, capture_fn=fake)
        self.assertEqual(out["blocked"], ["iaa"])
        self.assertEqual(out["lots"], 20)
        statuses = {r["url"]: r["last_status"] for r in db.query("SELECT url, last_status FROM sources")}
        self.assertEqual(statuses["https://www.iaai.com/blocked"], "blocked")
        self.assertIsNone(statuses["https://www.iaai.com/b"])  # skipped after the block

    def test_paste_parser(self):
        text = "2019 TOYOTA CAMRY SE\nLot # 70123456\nSeller: GEICO INSURANCE\nPrimary Damage: FRONT END\nLocation: TX - HOUSTON"
        out = capture.parse_listing_text(text)
        self.assertEqual((out["lot_id"], out["seller_raw"], out["year"], out["make"], out["state"], out["carrier_guess"]),
                         ("70123456", "GEICO INSURANCE", "2019", "TOYOTA", "TX", "GEICO"))

    def test_yard_directory_is_not_lots(self):
        # Copart's locations page lists yards: zip codes look like lot numbers, names contain a state and
        # 4-digit zip+4 suffixes look like years. None of it is a vehicle.
        yards = [{"facilityName": f"AL - YARD {i}", "zip1": f"350{i:02d}", "facilityStateCode": "AL",
                  "facilityPhoneNumber": f"205555{i:04d}", "facilityMailingZip2": str(1990 + i)} for i in range(18)]
        payloads = [{"url": "loc", "page_url": "https://www.copart.com/locations",
                     "data": {"data": {"locationByContinents": {"NORTH_AMERICA": {"AL": yards}}}}},
                    {"url": "api", "page_url": "https://www.copart.com/s?st=TX", "data": copart_like_payload(12)}]
        rows = capture.extract(payloads, "copart")
        self.assertEqual(len(rows), 12)
        self.assertEqual({r["source_url"] for r in rows}, {"https://www.copart.com/s?st=TX"})

    def test_raw_captures_are_never_overwritten(self):
        first, second = capture.save_raw([{"n": 1}], "copart_recorded"), capture.save_raw([{"n": 2}], "copart_recorded")
        self.assertNotEqual(first, second)
        self.assertEqual(len(list((paths.DATA / "raw").glob("*copart_recorded*.json.gz"))), 2)

    def test_dated_sale_list_is_not_learned(self):
        url = "https://www.copart.com/saleListResult/366/2026-09-28?location=FL+-+Clewiston"
        rows = capture.extract([{"url": "api", "page_url": url, "data": copart_like_payload(12)}], "copart")
        self.assertEqual(len(rows), 12)
        self.assertEqual(capture.learn_sources(rows, "copart", min_lots=10), [])  # that sale is over by next week

    def test_recording_survives_bad_data_and_a_crash(self):
        good = {"url": "api", "page_url": "https://www.copart.com/s?st=TX", "data": copart_like_payload(12)}
        odd = {"url": "odd", "page_url": "https://www.copart.com/", "data": {"x": [{"a": 1}] * 3}}
        real_extract = capture.extract

        def flaky(items, *a, **k):
            if items[0]["url"] == "odd":
                raise ValueError("unexpected data")
            return real_extract(items, *a, **k)

        def session(start_url, max_minutes, stop_file, on_payload, show_browser, script):
            on_payload(odd)
            on_payload(good)
            raise RuntimeError("browser crashed")

        said = []
        with mock.patch.object(capture, "extract", flaky), mock.patch.object(capture, "record_session", session):
            with self.assertRaises(RuntimeError):
                capture.run_guided(settings.load(), "copart", progress=lambda msg, **k: said.append(msg))
        self.assertEqual(db.scalar("SELECT COUNT(*) FROM observations"), 12)  # lots seen before the crash are kept
        self.assertEqual(len(list((paths.DATA / "raw").glob("*_copart_recorded.json.gz"))), 1)
        self.assertTrue(any("unexpected data" in m for m in said))

    def test_silent_imperva_block_page(self):
        normal = '<html><head><script src="/_Incapsula_Resource?SWJIYLWA=719d"></script></head><body>Results</body></html>'
        self.assertTrue(capture.looks_blocked("", "", BLOCK_PAGE))  # the message sits inside a frame; page text is empty
        self.assertFalse(capture.looks_blocked("Search results", "Results", normal))  # every page carries Imperva's script
        db.upsert("sources", [{"platform": "iaa", "url": "https://www.iaai.com/a", "label": "TX", "enabled": 1}], ["url"])
        s = settings.update(pause_between_pages=[0, 0], respect_robots_txt=False)
        out = capture.run_auto(s, progress=lambda *a, **k: None,
                               capture_fn=lambda *a: {"payloads": [], "title": "", "text": "", "html": BLOCK_PAGE})
        self.assertEqual(out["blocked"], ["iaa"])


BLOCK_PAGE = ('<html><head><meta name="ROBOTS" content="NOINDEX, NOFOLLOW"></head><body><iframe id="main-iframe" '
              'src="/_Incapsula_Resource?CWUDNSAI=23&incident_id=1-2">Request unsuccessful. Incapsula incident ID: 1-2'
              '</iframe></body></html>')
COPART_ROBOTS = "User-agent: *\nDisallow: /public/data/\nAllow: /lotSearchResults$\nDisallow: /lotSearchResults/\n"


class RobotsTests(TempData, unittest.TestCase):
    def setUp(self):
        super().setUp()
        net.forget_robots()

    def test_browser_captures_read_robots_in_the_browser(self):
        seen = []

        def browser(url):
            seen.append(url)
            return 200, COPART_ROBOTS

        # Copart's firewall refuses plain requests, and each refusal counts against the visitor: never send one
        with mock.patch.object(net, "http_get", side_effect=AssertionError("plain request sent")):
            self.assertTrue(net.robots_allows("https://www.copart.com/lotSearchResults?free=true&query=", fetch=browser))
            self.assertFalse(net.robots_allows("https://www.copart.com/lotSearchResults/?query=x", fetch=browser))
            self.assertFalse(net.robots_allows("https://www.copart.com/public/data/x", fetch=browser))
        self.assertEqual(seen, ["https://www.copart.com/robots.txt"])  # read once, then remembered

    def test_a_bot_check_instead_of_robots_stops_the_site(self):
        with self.assertRaises(net.RobotsBlocked):
            net.robots_allows("https://www.iaai.com/x", fetch=lambda url: (200, BLOCK_PAGE))
        with mock.patch.object(net, "http_get", return_value=FakeResp(BLOCK_PAGE, status=403)) as get:
            with self.assertRaises(net.RobotsBlocked):
                net.robots_allows("https://www.copart.com/x")
            with self.assertRaises(net.RobotsBlocked):  # no second attempt against a block
                net.robots_allows("https://www.copart.com/y")
        self.assertEqual(get.call_count, 1)

    def test_missing_robots_allows_and_server_error_disallows(self):
        with mock.patch.object(net, "http_get", return_value=FakeResp("not found", status=404)):
            self.assertTrue(net.robots_allows("https://a.example/x"))
        with mock.patch.object(net, "http_get", return_value=FakeResp("oops", status=503)):
            self.assertFalse(net.robots_allows("https://b.example/x"))

    def test_run_auto_carries_on_when_a_browser_fails(self):
        db.upsert("sources", [{"platform": "copart", "url": "https://www.copart.com/a", "label": "TX", "enabled": 1},
                              {"platform": "iaa", "url": "https://www.iaai.com/b", "label": "TX", "enabled": 1}], ["url"])
        s = settings.update(pause_between_pages=[0, 0], respect_robots_txt=True)

        def robots(url, **k):
            if "copart" in url:
                raise RuntimeError("Page.goto: Timeout 60000ms exceeded")
            return True

        fake = lambda url, *a: {"payloads": [{"url": "api", "page_url": url, "data": iaa_like_payload(12)}],  # noqa: E731
                                "title": "ok", "text": ""}
        with mock.patch.object(net, "robots_allows", side_effect=robots):
            out = capture.run_auto(s, progress=lambda *a, **k: None, capture_fn=fake)
        self.assertEqual(out["lots"], 12)  # IAA still captured
        status = {r["platform"]: r["last_status"] for r in db.query("SELECT platform, last_status FROM sources")}
        self.assertTrue(status["copart"].startswith("error"), status)

    def test_run_auto_stops_a_platform_after_two_empty_pages(self):
        # a block page that no marker recognises must not be hammered page after page
        db.upsert("sources", [{"platform": "copart", "url": f"https://www.copart.com/{i}", "label": "TX", "enabled": 1}
                              for i in range(4)], ["url"])
        s = settings.update(pause_between_pages=[0, 0], respect_robots_txt=False)
        calls = []

        def empty(url, *a):
            calls.append(url)
            return {"payloads": [], "title": "Just a moment", "text": "", "html": ""}

        out = capture.run_auto(s, progress=lambda *a, **k: None, capture_fn=empty)
        self.assertEqual((len(calls), out["blocked"]), (2, ["copart"]))

    def test_run_auto_stops_a_platform_when_robots_is_a_bot_check(self):
        db.upsert("sources", [{"platform": "copart", "url": "https://www.copart.com/a", "label": "TX", "enabled": 1},
                              {"platform": "copart", "url": "https://www.copart.com/b", "label": "FL", "enabled": 1}], ["url"])
        s = settings.update(pause_between_pages=[0, 0], respect_robots_txt=True)
        calls = []
        with mock.patch.object(net, "robots_allows", side_effect=net.RobotsBlocked("bot check")):
            out = capture.run_auto(s, progress=lambda *a, **k: None, capture_fn=lambda *a: calls.append(a))
        self.assertEqual((out["blocked"], calls), (["copart"], []))


class CliTests(TempData, unittest.TestCase):
    def run_cli(self, *argv):
        import cli
        with mock.patch.object(sys, "argv", ["cli.py", *argv]), mock.patch("builtins.print"), \
                mock.patch("traceback.print_exc"):
            return cli.main()

    def test_single_updates_record_freshness(self):
        with mock.patch.object(news, "run", return_value={"articles": 3, "new": 1}):
            self.assertEqual(self.run_cli("update", "news"), 0)
        with mock.patch.object(macro, "run", side_effect=RuntimeError("FRED is down")):
            self.assertEqual(self.run_cli("update", "macro"), 1)
        with mock.patch.object(capture, "run_auto", return_value={"pages": 1, "lots": 5, "blocked": []}):
            self.assertEqual(self.run_cli("capture", "auto"), 0)
        runs = db.last_runs()
        self.assertEqual({k: runs[k]["status"] for k in ("news", "macro", "capture")},
                         {"news": "ok", "macro": "failed", "capture": "ok"})
        self.assertIsNotNone(db.last_success("news"))


class TrackerTests(TempData, unittest.TestCase):
    def setUp(self):
        super().setUp()
        rows = []
        for wk, base in (("2026-09-15", 0), ("2026-09-22", 100000)):
            for st in ("TX", "FL"):
                o = base + (0 if st == "TX" else 50000)
                rows += split(wk, st, "GEICO INSURANCE", 19, 1, o)
                rows += split(wk, st, "PROGRESSIVE CASUALTY", 1, 19, o + 1000)
                sf_c, sf_i = (15, 5) if wk == "2026-09-15" else (8, 12)
                rows += split(wk, st, "STATE FARM MUTUAL", sf_c, sf_i, o + 2000)
            rows += split(wk, "CA", "STATE FARM MUTUAL", 30, 0, base + 90000)  # CA: Copart only -> unpaired
            rows += split(wk, "TX", "ACME RECIPROCAL EXCH", 3, 3, base + 95000)
        db.save_observations(rows)
        self.s = settings.load()

    def test_metrics_pairing_and_ci(self):
        m = tracker.metrics(self.s)
        self.assertEqual(m["week"], "2026-09-21")
        self.assertEqual(m["paired_states"], ["FL", "TX"])
        self.assertEqual(m["unpaired_states"], ["CA"])
        g = m["carriers"]["GEICO"]
        self.assertEqual((g["copart"], g["iaa"], g["n"]), (38, 2, 40))
        self.assertAlmostEqual(g["share"], 0.95)
        self.assertLess(g["lo"], 0.95)
        self.assertEqual(m["carriers"]["State Farm"]["n"], 40)  # CA's 30 unpaired lots excluded
        self.assertEqual(m["overall"]["n"], 126)

    def test_history_state_flow_and_overrides(self):
        h = tracker.weekly_history(self.s)
        self.assertEqual(sorted(h["week"].unique()), ["2026-09-14", "2026-09-21"])
        self.assertIn("All insurers", set(h["carrier"]))
        self.assertEqual(len(tracker.state_table(self.s, "2026-09-21")), 7)  # TX: 4 groups, FL: 3
        flow = tracker.new_lot_flow(self.s)
        self.assertEqual(set(flow["week"]), {"2026-09-21"})
        um = tracker.unmapped_sellers(self.s)
        self.assertEqual(um.iloc[0]["seller_raw"], "ACME RECIPROCAL EXCH")
        s2 = settings.update(carrier_overrides={"ACME RECIPROCAL EXCH": "Travelers"})
        self.assertIn("Travelers", tracker.metrics(s2)["carriers"])

    def test_one_site_only_does_not_crash(self):
        # the owner's data on Sep 27: saved Copart pages replayed, no IAA lots, so no state is paired
        db.execute("DELETE FROM observations")
        db.save_observations(lots("copart", "GA", "Dealer", 20, 0, "2026-09-27")
                             + lots("copart", "FL", "STATE FARM INSURANCE", 20, 100, "2026-09-27"))
        m = tracker.metrics(self.s)
        self.assertEqual((m["paired_states"], m["overall"]["n"], m["unpaired_states"]), ([], 0, ["FL", "GA"]))
        self.assertIn("only one site", tracker.split_problem(m))
        cur, changes = weekly.current(self.s)
        self.assertTrue(any("Insurer split paused" in c["text"] for c in changes))
        self.assertTrue(tracker.weekly_history(self.s).empty)

    def test_two_prop(self):
        self.assertAlmostEqual(tracker.two_prop_p(50, 100, 50, 100), 1.0)
        self.assertAlmostEqual(tracker.two_prop_p(60, 100, 40, 100), 0.0047, 3)

    def test_what_changed(self):
        cur, changes = weekly.current(self.s)  # compares with the previous week of listing data automatically
        first = weekly.what_changed(self.s, cur, None)
        self.assertTrue(any("first week of listing data" in c["text"] for c in first))
        sf = [c for c in changes if c["text"].startswith("State Farm: Copart share fell")]
        self.assertEqual(len(sf), 1)
        self.assertEqual(sf[0]["level"], "alert")  # 75% -> 40% on n=40 is significant
        self.assertTrue(any(c["level"] == "good" and c["text"].startswith("GEICO") for c in changes))
        self.assertTrue(any("CA" in c["text"] and c["level"] == "action" for c in changes))
        self.assertEqual(changes[0]["level"], "alert")  # alerts sort first


def fred_csv(sid, start, growth, header="observation_date"):
    dates = pd.date_range("2014-01-01", "2026-08-01", freq="MS")
    vals = [start * (1 + growth) ** (i / 12) for i in range(len(dates))]
    rows = [f"{header},{sid}"] + [f"{d.date()},{v:.3f}" for d, v in zip(dates, vals)]
    rows[4] = rows[4].split(",")[0] + ",."
    return "\n".join(rows)


def bls_json(sid, start, growth, first_year=2017, last=(2026, 8)):
    months = [(y, m) for y in range(first_year, last[0] + 1) for m in range(1, 13) if (y, m) <= last]
    data = [{"year": str(y), "period": f"M{m:02d}", "periodName": "", "value": f"{start * (1 + growth) ** (i / 12):.3f}"}
            for i, (y, m) in enumerate(months)]
    data.append({"year": "2025", "period": "M13", "periodName": "Annual", "value": "1.0"})  # annual averages are skipped
    return {"status": "REQUEST_SUCCEEDED", "message": [], "Results": {"series": [{"seriesID": sid, "data": data[::-1]}]}}


class CollectorTests(TempData, unittest.TestCase):
    def test_news_new_counts(self):
        now = pd.Timestamp.now(tz="UTC")
        rss = f"""<rss><channel>
          <item><title>Copart loses share as IAA volumes jump, data shows - Wire</title><link>https://e.com/a</link>
          <pubDate>{now:%a, %d %b %Y %H:%M:%S} GMT</pubDate><source url="x">Wire</source>
          <description>web-scraping data tracking units</description></item>
          <item><title>Copart loses share as IAA volumes jump, data shows - Other</title><link>https://e.com/b</link>
          <pubDate>{now:%a, %d %b %Y %H:%M:%S} GMT</pubDate></item>
          <item><title>Ancient - Old</title><link>https://e.com/c</link><pubDate>Mon, 01 Jan 2018 10:00:00 GMT</pubDate></item>
        </channel></rss>"""
        s = settings.update(news_queries=["q1", "q2"], extra_feeds=[])
        with mock.patch.object(net, "http_get", return_value=FakeResp(rss)):
            first = news.run(s, progress=lambda *a, **k: None)
            second = news.run(s, progress=lambda *a, **k: None)
        self.assertEqual((first["articles"], first["new"], second["new"]), (1, 1, 0))
        row = db.query("SELECT * FROM news")[0]
        self.assertIn("q1 | q2", row["queries"])
        self.assertGreaterEqual(row["score"], 4)

    def test_macro(self):
        csvs = {"CUSR0000SETD": fred_csv("CUSR0000SETD", 280, 0.05), "CUSR0000SETC": fred_csv("CUSR0000SETC", 140, 0.04, "DATE"),
                "CUSR0000SETA02": fred_csv("CUSR0000SETA02", 140, 0.0), "M12MTVUSM227NFWA": fred_csv("M12MTVUSM227NFWA", 3e6, 0.01)}
        posted = []

        def bls(url, payload, **k):
            posted.append(payload)
            return FakeResp(obj=bls_json("CUUR0000SETE", 450, 0.03))

        with mock.patch.object(net, "http_get", side_effect=lambda url, **k: FakeResp(csvs[url.split("id=")[1]])), \
                mock.patch.object(net, "http_post", side_effect=bls):
            latest = macro.run(settings.load(), progress=lambda *a, **k: None)
        self.assertGreater(latest["pressure_repair"]["value"], 130)
        self.assertAlmostEqual(latest["repair_minus_used"]["value"], 5.0, 1)
        self.assertIn("year_ago", latest["pressure_repair"])
        self.assertAlmostEqual(macro.frame(settings.load())["pressure_repair"].loc["2019"].mean(), 100, 1)
        self.assertAlmostEqual(latest["insurance_cpi_yoy"]["value"], 3.0, 1)  # insurance prices now come from BLS
        self.assertEqual(posted[0]["seriesid"], ["CUUR0000SETE"])


# Laid out like the real monthly releases (checked Sep 2026): whole thousands and growth without a % sign.
# The premiums table with its own "Direct - auto" row comes first, as a decoy.
PGR_EXHIBIT = """<html><body><p>The Progressive Corporation Reports August Results</p>
<table><tr><td>Direct &ndash; auto</td><td>$ 2,345.6</td><td>$ 2,100.0</td><td>12 %</td></tr>
<tr><td>Combined ratio</td><td>89.3</td><td>83.1</td><td>6.2 pts.</td></tr></table>
<p>August 31, (thousands; unaudited) 2026 2025 % Change</p><p>Policies in Force</p><table>
<tr><td>Agency &ndash; auto</td><td>11,343</td><td>10,575</td><td>7</td></tr>
<tr><td>Direct &ndash; auto</td><td>16,879</td><td>15,524</td><td>9</td></tr></table></body></html>"""
# Shaped like Copart's filings (checked Sep 2026). A 10-K first mentions the keyword in its risk factors, then
# carries the same notes paragraph as the 10-Qs near the end, split by a page break. Each notes paragraph is
# followed by the next note, whose number changes between filings.
DOJ_K = ("<p>Regulation. We must comply with laws relating to anti-money laundering and exporting. As described under "
         "Note 15 — Commitments and Contingencies, the U.S. Department of Justice is conducting an investigation "
         "into potential violations of certain money laundering laws. Changes in laws may affect us.</p>"
         "<p>NOTE 15 — Commitments and Contingencies. The Company maintains insurance. The U.S. Department of "
         "Justice (DOJ) is conducting an ongoing investigation into potential violations of certain money laundering "
         "laws. The Company received a letter from the DOJ in October 2023. We are unable to predict the range of "
         "81 Table of Contents possible loss. NOTE 16 — Guarantees. The Company guarantees some leases.</p>")
DOJ_Q1 = ("<p>The Company maintains insurance. The U.S. Department of Justice (DOJ) is conducting an ongoing investigation "
          "into potential violations of certain money laundering laws. The Company received a letter from the DOJ in "
          "October 2023. As of January 31, 2026 we are unable to predict the range of possible loss. NOTE 10 – "
          "Segments. The Company's regions are two segments.</p>")
# Only an as-of date, quote marks, a hyphen and the next note's number differ: not a wording change
DOJ_Q2 = (DOJ_Q1.replace("January 31, 2026", "April 30, 2026").replace("NOTE 10", "NOTE 11")
          .replace("(DOJ)", '("DOJ")').replace("money laundering laws", "money-laundering laws"))
DOJ_Q3 = DOJ_Q1.replace("As of January 31, 2026 we are unable to predict the range of possible loss.",
                        "The Company reached an agreement in principle with the DOJ and recorded an accrual of $50 million.")


class SecTests(TempData, unittest.TestCase):
    def setUp(self):
        super().setUp()
        sec._cik.clear()

    def fake(self, url, params=None, headers=None, **kw):
        assert headers and "@" in headers["User-Agent"]
        if url == sec.TICKERS_URL:
            return FakeResp(obj={"0": {"cik_str": 111, "ticker": "CPRT"}, "1": {"cik_str": 222, "ticker": "PGR"}})
        if "CIK0000000222" in url:
            return FakeResp(obj={"filings": {"recent": {"form": ["8-K", "10-Q"], "filingDate": ["2026-09-16", "2026-08-01"],
                                                        "accessionNumber": ["0000222-26-000009", "x"], "primaryDocument": ["a.htm", "b.htm"]}}})
        if "CIK0000000111" in url:
            return FakeResp(obj={"filings": {"recent": {
                "form": ["10-K", "10-Q", "10-Q", "10-Q", "10-K"],
                "filingDate": ["2026-09-25", "2026-05-29", "2026-03-03", "2025-11-24", "2025-09-26"],
                "accessionNumber": ["a-5", "a-4", "a-3", "a-2", "a-1"],
                "primaryDocument": ["k2.htm", "q3.htm", "q2.htm", "q1.htm", "k1.htm"]}}})
        if url.endswith("index.json"):
            return FakeResp(obj={"directory": {"item": [{"name": "pgr-ex99.htm"}, {"name": "logo.jpg"}]}})
        if url.endswith("pgr-ex99.htm"):
            return FakeResp(PGR_EXHIBIT)
        if "efts.sec.gov" in url:
            return FakeResp(obj={"hits": {"hits": [{"_id": "0000999-26-000001:rba10q.htm", "_source": {
                "ciks": ["0000000999"], "display_names": ["RB Global (RBA)"], "file_date": "2026-08-05", "form": "10-Q"}}]}})
        docs = {"k1.htm": DOJ_K, "q1.htm": DOJ_Q1, "q2.htm": DOJ_Q2, "q3.htm": DOJ_Q3, "k2.htm": DOJ_K}
        return FakeResp(f"<html><body>{docs[url.rsplit('/', 1)[-1]]}</body></html>")

    def test_everything(self):
        s = settings.update(user_name="Test Student", user_email="t@school.edu",
                            sec_release_targets=[{"ticker": "PGR", "filings": 1, "keywords": ["combined ratio"]}],
                            sec_fulltext_queries=['"automotive pricing incentives"'], doj_filings=5)
        with mock.patch.object(net, "http_get", side_effect=self.fake):
            sec.run(s, progress=lambda *a, **k: None)
        pgr = db.query("SELECT * FROM pgr_monthly")[0]
        self.assertEqual((pgr["period"], pgr["combined_ratio"], pgr["combined_ratio_prior"]), ("August 2026", 89.3, 83.1))
        self.assertEqual((pgr["direct_auto_pif"], pgr["direct_auto_growth"], pgr["agency_auto_growth"]), (16879.0, 9.0, 7.0))
        self.assertTrue(db.query("SELECT url FROM fulltext_hits")[0]["url"].endswith("/999/000099926000001/rba10q.htm"))
        # 10-Qs are compared with 10-Qs and 10-Ks with 10-Ks; date-only edits aren't changes
        changed = [(r["form"], r["changed"]) for r in db.query("SELECT form, changed FROM doj ORDER BY filing_date")]
        self.assertEqual(changed, [("10-K", None), ("10-Q", None), ("10-Q", 0), ("10-Q", 1), ("10-K", 0)])
        q = db.query("SELECT disclosure FROM doj WHERE filing_date='2026-03-03'")[0]["disclosure"]
        self.assertTrue(q.startswith("The U.S. Department of Justice"), q)
        self.assertNotIn("Segments", q)
        k = db.query("SELECT disclosure FROM doj WHERE filing_date='2025-09-26'")[0]["disclosure"]
        self.assertTrue(k.startswith("The U.S. Department of Justice"), k)  # the notes paragraph, not the risk factor
        self.assertTrue(k.endswith("range of possible loss."), k)  # page break removed, next note cut off
        st = sec.doj_status()
        self.assertFalse(st["changed_latest"])
        self.assertEqual(st["last_change"], "2026-05-29")

    def test_short_doj_passage_stops_at_the_next_note(self):
        # once the matter is resolved the paragraph may be one sentence; the next note must not leak in
        text = ("The Company maintains insurance. The U.S. Department of Justice closed its money laundering review. "
                "NOTE 11 — Segments. The Company has two segments.")
        self.assertEqual(sec.doj_passage(text, "money laundering"),
                         "The U.S. Department of Justice closed its money laundering review.")

    def test_progressive_period_year_and_negative_growth(self):
        out = sec.parse_progressive("The Progressive Corporation Reports December Results. Policies in Force "
                                    "Direct – auto 16,000 16,320 (2) %", "2026-01-28")
        self.assertEqual((out["period"], out["direct_auto_pif"], out["direct_auto_growth"]), ("December 2025", 16000.0, -2.0))

    def test_needs_user_details(self):
        with self.assertRaises(sec.SecConfigError):
            sec.run(settings.load())

    def test_word_diff(self):
        out = sec.word_diff("we cannot estimate a range", "we recorded a $50 million accrual")
        self.assertIn("<del>cannot estimate</del>", out)
        self.assertIn("<ins>recorded</ins>", out)


class WeeklyReportTests(TempData, unittest.TestCase):
    def test_weekly_run_report_excel(self):
        rows = []
        for st in ("TX", "FL"):
            rows += split("2026-09-22", st, "GEICO INSURANCE", 19, 1, 0 if st == "TX" else 5000)
            rows += split("2026-09-22", st, "PROGRESSIVE CASUALTY", 1, 19, 1000 if st == "TX" else 6000)
        db.save_observations(rows)
        s = settings.load()
        ok = lambda *a, **k: {"ok": True}  # noqa: E731
        with mock.patch.object(news, "run", ok), mock.patch.object(macro, "run", ok), \
                mock.patch.object(capture, "run_scan", ok), \
                mock.patch.object(sec, "run", side_effect=sec.SecConfigError("add details")):
            out = weekly.run(s, progress=lambda *a, **k: None)
        self.assertEqual(out["results"], {"news": "ok", "macro": "ok", "sec": "skipped", "scan": "ok"})
        html = (paths.DATA / "reports" / out["report"]).read_text()
        self.assertIn("Where each insurer", html)
        self.assertIn("data:image/png;base64,", html)
        self.assertEqual(db.last_runs()["weekly"]["status"], "ok")
        self.assertEqual(len(report.list_reports()), 1)
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(report.excel(s)))
        self.assertIn("Insurer split (latest)", wb.sheetnames)
        self.assertIn("All lots", wb.sheetnames)


class WeeklyFailureTests(TempData, unittest.TestCase):
    def test_a_failed_analysis_step_is_recorded(self):
        ok = lambda *a, **k: {"ok": True}  # noqa: E731
        with mock.patch.object(news, "run", ok), mock.patch.object(macro, "run", ok), mock.patch.object(sec, "run", ok), \
                mock.patch.object(capture, "run_scan", ok), mock.patch.object(weekly, "current", side_effect=KeyError("platform")):
            with self.assertRaises(KeyError):
                weekly.run(settings.load(), progress=lambda *a, **k: None)
        runs = db.last_runs()
        self.assertEqual((runs["weekly"]["status"], runs["news"]["status"]), ("failed", "ok"))  # data steps still count
        self.assertIn("platform", runs["weekly"]["message"])


class SchedulerJobTests(TempData, unittest.TestCase):
    def test_plist_and_warning(self):
        s = settings.load()
        pl = scheduler.build_plist(s)
        self.assertEqual(pl["ProgramArguments"][-1], "scheduled")
        self.assertEqual(pl["StartCalendarInterval"], {"Hour": 8, "Minute": 0})  # daily; the full update stays weekly
        from pathlib import Path
        with mock.patch.object(paths, "ROOT", Path.home() / "Downloads" / "cprt-console"):
            self.assertIn("Downloads", scheduler.protected_folder_warning())
        with mock.patch.object(scheduler, "supported", return_value=False):
            with self.assertRaises(RuntimeError):
                scheduler.install(s)

    def test_background_job_lifecycle(self):
        job = jobs.start("report", ["report"])
        self.assertEqual(job["status"], "running")
        with self.assertRaises(RuntimeError):
            jobs.start("report", ["report"])  # one at a time
        for _ in range(120):
            if not jobs.running():
                break
            time.sleep(0.5)
        jobs._procs[job["id"]].wait(timeout=30)
        final = jobs.latest()
        self.assertEqual(final["status"], "done", final)
        self.assertTrue(final["log_tail"])


if __name__ == "__main__":
    unittest.main(verbosity=2)
